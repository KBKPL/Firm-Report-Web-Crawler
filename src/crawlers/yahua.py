import os
import logging
import urllib.parse
from typing import List, Dict, Optional

from bs4 import BeautifulSoup
from docx import Document

from src.crawlers.base import CompanyCrawler
from src.utils.http_utils import session
from src.utils.pdf_utils import extract_text_from_pdf
from src.utils.text_utils import sanitize_text, find_paragraphs_with_keyword
from src.utils.docx_utils import add_keyword_paragraphs, save_crawler_docx
from src.utils.html_utils import fetch_rendered_html, find_paragraphs_from_html

logger = logging.getLogger(__name__)

class YahuaCrawler(CompanyCrawler):
    SECTIONS = {
        "1": ("业绩报告", "quarterly performance", "crawl_quarterly_performance"),
        "2": ("投资者交流", "investor communication", "crawl_company_announcements"),
        "3": ("券商研报", "broker reports", "crawl_broker_reports"),
        "4": ("企业动态", "company news", "crawl_company_news"),
    }

    def __init__(self, full_code: str, config: dict):
        super().__init__(full_code, config)
        self.base_url = config.get("base_url")
        self.page_size = config.get("page_size", 10)
        self.quarterly_url = config.get("quarterly_url")
        self.company_announcement_url = config.get("company_announcement_url")
        self.broker_report_url = config.get("broker_report_url")
        self.company_news_url = config.get("company_news_url")
        self.company_news_page_size = config.get("company_news_page_size", 4)

    def fetch_quarterly_performance_page(self, page_index: int) -> List[Dict[str, str]]:
        """Fetch one page of Yahua quarterly performance metadata."""
        page_num = page_index + 1
        url = self.quarterly_url.format(page=page_num)
        resp = session.get(url, timeout=10)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        items = soup.find_all("li", class_="clearfix")
        records: List[Dict[str, str]] = []
        for item in items:
            title_el = item.find("span", class_="ivtms eT")
            date_el = item.find("b", class_="ivldate")
            link = item.find("a", class_="linkA", href=True)
            if not title_el or not date_el or not link:
                logger.warning("Missing element in Yahua quarterly item; skipping.")
                continue
            raw_title = title_el.get_text(strip=True)
            title = raw_title.replace("雅化集团：", "", 1)
            pub_date = date_el.get_text(strip=True)
            partial = link["href"]
            pdf_url = urllib.parse.urljoin(self.base_url, partial)
            records.append({"title": title, "publishDate": pub_date, "pdf_url": pdf_url})
        return records

    def crawl_quarterly_performance(
        self,
        keywords: List[str],
        output_dir: str,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
    ) -> Dict[str, str]:
        os.makedirs(output_dir, exist_ok=True)
        docs = {kw: Document() for kw in keywords}
        generated: Dict[str, str] = {}
        page_index = 0
        while True:
            logger.info(f"Fetching Yahua quarterly page: {page_index}")
            records = self.fetch_quarterly_performance_page(page_index)
            if not records:
                break
            break_page = False
            for rec in records:
                title = rec["title"]
                pub_date = rec["publishDate"]
                if end_date and pub_date > end_date:
                    continue
                if start_date and pub_date < start_date:
                    break_page = True
                    break
                pdf_url = rec.get("pdf_url")
                if not pdf_url:
                    logger.warning(f"No PDF URL for record {title}; skipping.")
                    continue
                text = extract_text_from_pdf(pdf_url)
                text = sanitize_text(text)
                for kw in keywords:
                    paras = find_paragraphs_with_keyword(text, kw)
                    if not paras:
                        continue
                    doc = docs[kw]
                    doc.add_heading(f"{pub_date}_{title}", level=1)
                    add_keyword_paragraphs(doc, paras, kw, pdf_url)
            if break_page or len(records) < self.page_size:
                break
            page_index += 1
        section_label = self.SECTIONS["1"][0]
        for kw, doc in docs.items():
            path = save_crawler_docx(doc, self.full_code, kw, section_label, output_dir)
            generated[kw] = path
        return generated

    def fetch_company_announcements_page(self, page_index: int) -> List[Dict[str, str]]:
        """Fetch one page of Yahua investor communication metadata."""
        page_num = page_index + 1
        url = self.company_announcement_url.format(page=page_num)
        resp = session.get(url, timeout=10)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        items = soup.find_all("li", class_="clearfix")
        records: List[Dict[str, str]] = []
        prefix = f"{self.full_code}雅化集团"
        for item in items:
            title_el = item.find("span", class_="ivtms eT")
            date_el = item.find("b", class_="ivldate")
            link = item.find("a", class_="linkA", href=True)
            if not title_el or not date_el or not link:
                logger.warning("Missing element in Yahua investor communication item; skipping.")
                continue
            raw_title = title_el.get_text(strip=True)
            if raw_title.startswith(prefix):
                title = raw_title[len(prefix):]
            else:
                title = raw_title
            pub_date = date_el.get_text(strip=True)
            pdf_url = urllib.parse.urljoin(self.base_url, link["href"])
            records.append({"title": title, "publishDate": pub_date, "pdf_url": pdf_url})
        return records

    def crawl_company_announcements(
        self,
        keywords: List[str],
        output_dir: str,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
    ) -> Dict[str, str]:
        os.makedirs(output_dir, exist_ok=True)
        docs = {kw: Document() for kw in keywords}
        generated: Dict[str, str] = {}
        page_index = 0
        while True:
            logger.info(f"Fetching Yahua investor communication page: {page_index}")
            records = self.fetch_company_announcements_page(page_index)
            if not records:
                break
            break_page = False
            for rec in records:
                title = rec["title"]
                pub_date = rec["publishDate"]
                if end_date and pub_date > end_date:
                    continue
                if start_date and pub_date < start_date:
                    break_page = True
                    break
                pdf_url = rec.get("pdf_url")
                if not pdf_url:
                    logger.warning(f"No PDF URL for record {title}; skipping.")
                    continue
                text = extract_text_from_pdf(pdf_url)
                text = sanitize_text(text)
                for kw in keywords:
                    paras = find_paragraphs_with_keyword(text, kw)
                    if not paras:
                        continue
                    doc = docs[kw]
                    doc.add_heading(f"{pub_date}_{title}", level=1)
                    add_keyword_paragraphs(doc, paras, kw, pdf_url)
            if break_page or len(records) < self.page_size:
                break
            page_index += 1
        section_label = self.SECTIONS["2"][0]
        for kw, doc in docs.items():
            path = save_crawler_docx(doc, self.full_code, kw, section_label, output_dir)
            generated[kw] = path
        return generated

    def fetch_broker_reports_page(self, page_index: int) -> List[Dict[str, str]]:
        """Fetch one page of Yahua broker reports metadata."""
        page_num = page_index + 1
        url = self.broker_report_url.format(page=page_num)
        resp = session.get(url, timeout=10)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        items = soup.find_all("li", class_="clearfix")
        records: List[Dict[str, str]] = []
        for item in items:
            title_el = item.find("span", class_="ivtms eT")
            date_el = item.find("b", class_="ivldate")
            link = item.find("a", class_="linkA", href=True)
            if not title_el or not date_el or not link:
                logger.warning("Missing element in Yahua broker report item; skipping.")
                continue
            title = title_el.get_text(strip=True)
            pub_date = date_el.get_text(strip=True)
            pdf_url = urllib.parse.urljoin(self.base_url, link["href"])
            records.append({"title": title, "publishDate": pub_date, "pdf_url": pdf_url})
        return records

    def crawl_broker_reports(
        self,
        keywords: List[str],
        output_dir: str,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
    ) -> Dict[str, str]:
        os.makedirs(output_dir, exist_ok=True)
        docs = {kw: Document() for kw in keywords}
        generated: Dict[str, str] = {}
        page_index = 0
        while True:
            logger.info(f"Fetching Yahua broker reports page: {page_index}")
            records = self.fetch_broker_reports_page(page_index)
            if not records:
                break
            break_page = False
            for rec in records:
                title = rec["title"]
                pub_date = rec["publishDate"]
                if end_date and pub_date > end_date:
                    continue
                if start_date and pub_date < start_date:
                    break_page = True
                    break
                pdf_url = rec.get("pdf_url")
                if not pdf_url:
                    logger.warning(f"No PDF URL for record {title}; skipping.")
                    continue
                text = extract_text_from_pdf(pdf_url)
                text = sanitize_text(text)
                for kw in keywords:
                    paras = find_paragraphs_with_keyword(text, kw)
                    if not paras:
                        continue
                    doc = docs[kw]
                    doc.add_heading(f"{pub_date}_{title}", level=1)
                    add_keyword_paragraphs(doc, paras, kw, pdf_url)
            if break_page or len(records) < self.page_size:
                break
            page_index += 1
        section_label = self.SECTIONS["3"][0]
        for kw, doc in docs.items():
            path = save_crawler_docx(doc, self.full_code, kw, section_label, output_dir)
            generated[kw] = path
        return generated

    def fetch_company_news_page(self, page_index: int) -> List[Dict[str, str]]:
        """Fetch one page of Yahua company news metadata."""
        page_num = page_index + 1
        url = self.company_news_url.format(page=page_num)
        resp = session.get(url, timeout=10)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        records: List[Dict[str, str]] = []
        # iterate all linkA anchors to capture both layouts
        for link in soup.find_all("a", class_="linkA", href=True):
            parent = link.parent
            title_el = None
            date_el = None
            # layout 1: div.nwfirst
            if parent.name == "div" and "nwfirst" in parent.get("class", []):
                cont = parent.find("div", class_="nfcont")
                title_el = cont.find("h5", class_="nfctitle") if cont else None
                date_el = cont.find("span", class_="nfcdate") if cont else None
            # layout 2: li
            elif parent.name == "li":
                wz = parent.find("div", class_="nwlwz")
                title_el = wz.find("h5", class_="nwlbt") if wz else None
                date_el = wz.find("span", class_="nwldate") if wz else None
            else:
                continue
            if not title_el or not date_el:
                logger.warning("Missing elements in Yahua company news item; skipping.")
                continue
            raw_title = title_el.get_text(strip=True)
            raw_date = date_el.get_text(strip=True)
            # convert MM.DD.YYYY to YYYY-MM-DD
            parts = raw_date.split('.')
            if len(parts) == 3:
                month, day, year = parts
                pub_date = f"{year}-{month.zfill(2)}-{day.zfill(2)}"
            else:
                pub_date = raw_date
            full_url = urllib.parse.urljoin(self.base_url, link["href"])
            records.append({"title": raw_title, "publishDate": pub_date, "pdf_url": full_url})
        return records

    def crawl_company_news(
        self,
        keywords: List[str],
        output_dir: str,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
    ) -> Dict[str, str]:
        os.makedirs(output_dir, exist_ok=True)
        docs = {kw: Document() for kw in keywords}
        generated: Dict[str, str] = {}
        page_index = 0
        while True:
            logger.info(f"Fetching Yahua company news page: {page_index}")
            records = self.fetch_company_news_page(page_index)
            if not records:
                break
            break_page = False
            for rec in records:
                title = rec["title"]
                pub_date = rec["publishDate"]
                if end_date and pub_date > end_date:
                    continue
                if start_date and pub_date < start_date:
                    break_page = True
                    break
                url = rec.get("pdf_url")
                if not url:
                    logger.warning(f"No URL for company news {title}; skipping.")
                    continue
                # fetch text: PDF or HTML
                if url.lower().endswith(".pdf"):
                    text = extract_text_from_pdf(url)
                    url_used = url
                else:
                    html_text = fetch_rendered_html(url)
                    if not html_text:
                        continue
                    text = BeautifulSoup(html_text, "lxml").get_text("\n")
                    url_used = url
                text = sanitize_text(text)
                for kw in keywords:
                    paras = find_paragraphs_with_keyword(text, kw)
                    if not paras:
                        continue
                    doc = docs[kw]
                    doc.add_heading(f"{pub_date}_{title}", level=1)
                    add_keyword_paragraphs(doc, paras, kw, url_used)
            if break_page or len(records) < self.company_news_page_size:
                break
            page_index += 1
        section_label = self.SECTIONS["4"][0]
        for kw, doc in docs.items():
            path = save_crawler_docx(doc, self.full_code, kw, section_label, output_dir)
            generated[kw] = path
        return generated
